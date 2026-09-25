import json
from pathlib import Path
from zipfile import ZipFile

import pytest

from service_packets.finalize import finalize_packet, finalize_worker_export


def final_idml(path):
    with ZipFile(path, 'w') as archive:
        archive.writestr('mimetype', 'application/vnd.adobe.indesign-idml-package')
        archive.writestr('Resources/Styles.xml', '<Styles><ParagraphStyle Self="ParagraphStyle/OOW Body"/></Styles>')
        frames = []
        for number, label in enumerate(('service_heading', 'worship_order')):
            frames.append(f'<TextFrame ParentStory="s{number}"><Properties><Label><KeyValuePair Key="Label" Value="{label}"/></Label></Properties></TextFrame>')
            archive.writestr(f'Stories/s{number}.xml', f'<Document><Story Self="s{number}"><ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/OOW Body"><Content>Final approved text</Content></ParagraphStyleRange></Story></Document>')
        archive.writestr('Spreads/s.xml', '<Spread>'+''.join(frames)+'</Spread>')
    return path


def designation(path, service='main', day='2026-09-06'):
    return {'path': str(path), 'date': day, 'service': service,
            'user_designation': 'User explicitly selected this saved final bulletin.'}


def test_main_capture_is_immutable_and_early_final_does_not_replace_it(tmp_path):
    original = final_idml(tmp_path/'final.idml')
    root = tmp_path/'data'
    first = finalize_packet(root, designation(original))
    pointer = root/'state/service_packets/baseline.json'
    baseline = json.loads(pointer.read_text())
    assert baseline['service'] == 'main' and baseline['finalized'] is True
    assert Path(baseline['path']).read_bytes() == original.read_bytes()
    assert Path(baseline['path']) != original
    assert first['outstanding'] == ['cameras', 'sound']
    before = pointer.read_bytes()
    early = finalize_packet(root, designation(original, 'early_traditional'))
    assert early['promoted'] is False
    assert pointer.read_bytes() == before
    original.write_bytes(b'Later user edits')
    assert Path(baseline['path']).read_bytes() != original.read_bytes()


def test_supplied_final_technical_sheets_are_captured(tmp_path):
    from docx import Document
    source = final_idml(tmp_path/'final.idml')
    technical = tmp_path/'camera.docx'
    Document().save(technical)
    selected = designation(source)
    selected['technical'] = {'cameras': str(technical)}
    result = finalize_packet(tmp_path/'data', selected)
    assert result['outstanding'] == ['sound']
    assert Path(result['technical']['cameras']['path']).read_bytes() == technical.read_bytes()


def test_existing_finalization_lock_is_not_broken(tmp_path):
    source = final_idml(tmp_path/'final.idml')
    root = tmp_path/'data'
    state = root/'state/service_packets'
    state.mkdir(parents=True)
    lock = state/'finalize.lock'
    lock.write_text('Another process owns this lock')
    with pytest.raises(FileExistsError):
        finalize_packet(root, designation(source))
    assert lock.read_text() == 'Another process owns this lock'
    assert not (state/'baseline.json').exists()


@pytest.mark.parametrize('problem', ['empty', 'undefined_style', 'older', 'no_designation'])
def test_invalid_or_unapproved_main_keeps_existing_baseline(tmp_path, problem):
    source = final_idml(tmp_path/'final.idml')
    root = tmp_path/'data'
    finalize_packet(root, designation(source))
    pointer = root/'state/service_packets/baseline.json'
    before = pointer.read_bytes()
    selected = designation(source)
    if problem == 'older':
        selected['date'] = '2026-08-30'
    elif problem == 'no_designation':
        selected['user_designation'] = ''
    else:
        with ZipFile(source) as archive:
            members = {name: archive.read(name) for name in archive.namelist()}
        if problem == 'empty':
            members['Stories/s1.xml'] = members['Stories/s1.xml'].replace(b'Final approved text', b' ')
        else:
            members['Stories/s1.xml'] = members['Stories/s1.xml'].replace(b'ParagraphStyle/OOW Body', b'ParagraphStyle/Unknown')
        with ZipFile(source, 'w') as archive:
            for name, data in members.items():
                archive.writestr(name, data)
    with pytest.raises(ValueError):
        finalize_packet(root, selected)
    assert pointer.read_bytes() == before


def test_failed_atomic_promotion_preserves_previous_pointer(tmp_path, monkeypatch):
    source = final_idml(tmp_path/'final.idml')
    root = tmp_path/'data'
    finalize_packet(root, designation(source))
    pointer = root/'state/service_packets/baseline.json'
    before = pointer.read_bytes()
    def fail(*args):
        raise OSError('Injected storage failure')
    monkeypatch.setattr('service_packets.finalize.os.replace', fail)
    with pytest.raises(OSError, match='storage failure'):
        finalize_packet(root, designation(source, day='2026-09-13'))
    assert pointer.read_bytes() == before


@pytest.mark.parametrize('baseline_state', ['valid', 'stale', 'missing', 'modified'])
def test_main_stage_uses_previous_week_baseline_and_defers_early_services(tmp_path, baseline_state):
    from docx import Document
    from service_packets.build import build_packet
    from service_packets.model import ROLES
    source = final_idml(tmp_path/'final.idml')
    root = tmp_path/'data'
    final = finalize_packet(root, designation(source))
    if baseline_state == 'stale':
        pointer = root/'state/service_packets/baseline.json'
        record = json.loads(pointer.read_text())
        record['date'] = '2026-08-30'
        pointer.write_text(json.dumps(record))
    elif baseline_state == 'missing':
        Path(final['baseline']['path']).unlink()
    elif baseline_state == 'modified':
        Path(final['baseline']['path']).write_bytes(b'Changed since finalization')
    services = [{'key': role, 'time': time, 'title': 'Next week', 'order_source': 'order',
                 'items': [], 'heading_paragraphs': [{'style': 'OOW Body', 'text': time}],
                 'bulletin_paragraphs': [{'style': 'OOW Body', 'text': 'New order'}]}
                for role, time in zip(ROLES, ['08:30', '09:30', '10:30'])]
    packet = {'schema_version': 1, 'date': '2026-09-13', 'timezone': 'America/New_York',
              'schedule': {'kind': 'ordinary'}, 'expected_services': list(ROLES),
              'sources': [{'id': 'order', 'kind': 'nathan_docx', 'scope': list(ROLES)}],
              'services': services}
    templates = {}
    for kind, columns in [('cameras', 5), ('sound', 2)]:
        document = Document()
        document.add_paragraph('Header')
        document.add_table(rows=2, cols=columns)
        templates[kind] = tmp_path/(kind+'.docx')
        document.save(templates[kind])
    result = build_packet(packet, root, {'allowed_styles': ['OOW Body']}, {}, templates)
    if baseline_state != 'valid':
        assert len(result['artifacts']) == 3
        assert any(issue['code'] == 'baseline_unavailable' for issue in result['issues'])
        return
    assert len(result['artifacts']) == 4
    repeated = build_packet(packet, root, {'allowed_styles': ['OOW Body']}, {}, templates)
    assert repeated['revision'] == result['revision']
    assert repeated['reused'] is True
    corrected = build_packet(packet, root, {'allowed_styles': ['OOW Body']}, {'rule_change': True}, templates)
    assert corrected['revision'] != result['revision']
    edited = Path(corrected['artifacts'][0]['path'])
    edited.write_bytes(b'Human cleanup in progress')
    regenerated = build_packet(packet, root, {'allowed_styles': ['OOW Body']}, {'rule_change': True}, templates)
    assert regenerated['revision'] != corrected['revision']
    assert edited.read_bytes() == b'Human cleanup in progress'
    assert all(a['proof_status'] == 'pending' for a in repeated['artifacts'])
    saved = json.loads((Path(result['review_path']).parent/'service-packet.json').read_text())
    assert saved['baseline']['sha256'] == final['baseline']['sha256']
    for artifact in result['artifacts']:
        if artifact['format'] == 'idml':
            with ZipFile(artifact['path']) as output, ZipFile(source) as original:
                assert output.read('Spreads/s.xml') == original.read('Spreads/s.xml')


def test_indd_is_deferred_without_touching_source_or_promoting(tmp_path):
    source = tmp_path/'saved.indd'
    source.write_bytes(b'User document')
    result = finalize_packet(tmp_path/'data', designation(source))
    assert result == {'status': 'pending', 'reason': 'saved_idml_required', 'promoted': False}
    assert source.read_bytes() == b'User document'
    assert not (tmp_path/'data/state/service_packets/baseline.json').exists()


@pytest.mark.parametrize('busy', [True, False])
def test_indd_export_uses_only_scratch_and_defers_when_busy(tmp_path, busy):
    from hashlib import sha256
    source = tmp_path/'saved.indd'
    source.write_bytes(b'Saved user document')
    def adapter(request):
        assert Path(request['input']) != source
        assert Path(request['input']).read_bytes() == source.read_bytes()
        result = {'id': request['id'], 'input_hash': request['sha256']}
        if busy:
            return dict(result, proof_status='pending', reason='user_documents_open')
        exported = final_idml(Path(request['output']))
        return dict(result, proof_status='visual_review_pending', bad_fonts=0, bad_links=0,
                    output_hash=sha256(exported.read_bytes()).hexdigest())
    result = finalize_packet(tmp_path/'data', designation(source), adapter)
    assert result['promoted'] is (not busy)
    assert source.read_bytes() == b'Saved user document'
    if busy:
        assert result['reason'] == 'user_documents_open'
        assert not (tmp_path/'data/state/service_packets/baseline.json').exists()


@pytest.mark.parametrize('bad_links', [0, 1])
def test_registered_export_handoff_preserves_original_provenance(tmp_path, bad_links):
    from hashlib import sha256
    source=tmp_path/'final.indd'; source.write_bytes(b'approved original')
    exported=final_idml(tmp_path/'export.idml')
    request={'id':'export-1','operation':'export_final_idml','input_hash':sha256(source.read_bytes()).hexdigest(),
             'output':str(exported)}
    native={'id':'export-1','input_hash':request['input_hash'],'proof_status':'visual_review_pending',
            'output_hash':sha256(exported.read_bytes()).hexdigest(),'bad_fonts':0,'bad_links':bad_links}
    result={**{k:request[k] for k in ('id','operation','input_hash')},'status':'complete',
            'output_hash':native['output_hash'],'native':native}
    if bad_links:
        with pytest.raises(ValueError): finalize_worker_export(tmp_path/'data',designation(source),request,result)
        assert not (tmp_path/'data/state/service_packets/baseline.json').exists()
    else:
        record=finalize_worker_export(tmp_path/'data',designation(source),request,result)
        assert record['baseline']['original_source']['path']==str(source.resolve())
        assert record['baseline']['original_source']['sha256']==request['input_hash']
