from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from docx import Document
import pytest

from service_packets.bulletin import apply_bulletin_rules
from service_packets.build import build_packet
from service_packets.model import ROLES, validate_packet
from test_idml import baseline
from test_technical import templates


def packet(role='main', lines=None):
    lines = lines or [('sermon', 'Sermon')]
    evidence = [{'source_id': 'order', 'location': 'body/p/1'}]
    items = [{'id': str(i), 'kind': kind, 'source_wording': text,
              'bulletin_text': text, 'display_wording': text.split('\t')[0],
              'technical_wording': text.split('\t')[0], 'evidence': deepcopy(evidence),
              'proclaim': [{'label': 'Source', 'text': text, 'evidence': deepcopy(evidence)}]}
             for i, (kind, text) in enumerate(lines)]
    return {'schema_version': 1, 'date': '2099-09-20', 'timezone': 'America/New_York',
            'schedule': {'kind': 'ordinary'}, 'expected_services': list(ROLES),
            'sources': [{'id': 'order', 'kind': 'nathan_docx', 'scope': list(ROLES)}],
            'services': [{'key': role, 'time': {'main': '10:30', 'modern': '09:30',
                         'early_traditional': '08:30'}[role], 'order_source': 'order',
                         'items': items, 'heading_paragraphs': [{'style': 'Body', 'text': 'Sunday'}],
                         'bulletin_paragraphs': [{'style': 'Body', 'text': text} for _, text in lines]}],
            'issues': []}


@pytest.mark.parametrize('role', ROLES)
@pytest.mark.parametrize('kind,label,credit', [
    ('hymn_1', 'Hymn', 'UMH 95'), ('song', 'Song', 'Recording Artist'),
    ('anthem', 'Sanctuary Choir Anthem', 'Composer; arr. Arranger'),
    ('prelude', 'Prelude', 'Johann Crüger'), ('offertory', 'Offertory', 'Composer'),
    ('postlude', 'Postlude', 'Composer'), ('exit_music', 'Exit Music', 'Composer'),
    ('prayer_response', 'Prayer Response', 'Composer'),
    ('benediction_response', 'Choral Response', 'Composer'),
    ('communion_hymn', 'Hymn', 'TFWS 2223'), ('introit', 'Choral Introit', 'Composer'),
])
def test_quotes_only_selection_title_in_both_bulletin_fields(role, kind, label, credit):
    text = label + '\tGod’s Love — A / B\t' + credit + '\nVerse 1; Soloist'
    original = packet(role, [(kind, text)])
    before = deepcopy(original)
    result = apply_bulletin_rules(validate_packet(original))
    service = result['services'][0]
    expected = label + '\t“God’s Love — A / B”\t' + credit + '\nVerse 1; Soloist'
    assert service['items'][0]['bulletin_text'] == expected
    assert service['bulletin_paragraphs'][0]['text'] == expected
    for field in ('source_wording', 'display_wording', 'technical_wording', 'evidence', 'proclaim'):
        assert service['items'][0][field] == before['services'][0]['items'][0][field]
    assert original == before
    assert apply_bulletin_rules(deepcopy(result)) == result


@pytest.mark.parametrize('title', ['“Already Quoted”', '"Already Quoted"', '"“Already Quoted”"'])
def test_normalizes_existing_outer_double_quotes(title):
    result = apply_bulletin_rules(packet(lines=[('song', 'Song\t' + title + '\tArtist')]))
    assert result['services'][0]['bulletin_paragraphs'][0]['text'] == 'Song\t“Already Quoted”\tArtist'


@pytest.mark.parametrize('kind,text', [
    ('reading', 'First Reading\tJohn 13.1-5 (CEB)'), ('prayer', 'Prayer\tOur community'),
    ('doxology', 'Doxology\t\tUMH 95'), ('doxology', 'Doxology\tUMH 95'),
    ('prelude', 'Prelude'), ('song', 'Song\t\tArtist'),
    ('prelude', 'Prelude\tREVIEW: selection not supplied'),
])
def test_does_not_quote_non_titles_or_missing_selection_placeholders(kind, text):
    result = apply_bulletin_rules(packet(lines=[(kind, text)]))
    assert result['services'][0]['items'][0]['bulletin_text'] == text
    assert result['services'][0]['bulletin_paragraphs'][0]['text'] == text


@pytest.mark.parametrize('role', ROLES)
@pytest.mark.parametrize('speaker', [None, 'Rev. Guest Speaker'])
def test_sermon_default_and_evidenced_override_update_both_bulletin_fields(role, speaker):
    original = packet(role)
    item = original['services'][0]['items'][0]
    if speaker:
        item['values'] = {'speaker': {'status': 'supplied', 'value': speaker,
                                     'evidence': deepcopy(item['evidence'])}}
        # Accepted source corrections override earlier authored bulletin fields.
        item['bulletin_text'] = 'Sermon\t\tRev. Nathan Howe'
    before = deepcopy(original)
    result = apply_bulletin_rules(validate_packet(original))
    service = result['services'][0]
    expected = 'Sermon\t\t' + (speaker or 'Rev. Nathan Howe')
    assert service['items'][0]['bulletin_text'] == expected
    assert service['bulletin_paragraphs'][0]['text'] == expected
    assert service['items'][0]['technical_wording'] == 'Sermon'
    assert service['items'][0]['source_wording'] == 'Sermon'
    assert original == before


def test_retains_existing_guest_speaker_and_formats_paragraph_without_item_text():
    original = packet(lines=[('sermon', 'Sermon\t\tRev. Guest Speaker')])
    del original['services'][0]['items'][0]['bulletin_text']
    result = apply_bulletin_rules(validate_packet(original))
    service = result['services'][0]
    assert service['items'][0]['bulletin_text'] == 'Sermon\t\tRev. Guest Speaker'
    assert service['bulletin_paragraphs'][0]['text'] == 'Sermon\t\tRev. Guest Speaker'


@pytest.mark.parametrize('case', ['conflict', 'unresolved', 'middle_field', 'extra_field'])
def test_ambiguous_sermon_fields_are_preserved_and_flagged(case):
    original = packet()
    service = original['services'][0]
    if case == 'conflict':
        service['items'][0]['bulletin_text'] = 'Sermon\t\tRev. One'
        service['bulletin_paragraphs'][0]['text'] = 'Sermon\t\tRev. Two'
    elif case == 'unresolved':
        service['items'][0]['values'] = {'speaker': {'status': 'unresolved'}}
    else:
        text = 'Sermon\tPossible Title' if case == 'middle_field' else 'Sermon\t\tRev. Guest\tOther'
        service['items'][0]['bulletin_text'] = service['bulletin_paragraphs'][0]['text'] = text
    before = deepcopy(original)
    result = apply_bulletin_rules(validate_packet(original))
    assert result['services'] == before['services']
    assert any(i['code'] == 'sermon_speaker_review' for i in result['issues'])
    assert apply_bulletin_rules(deepcopy(result)) == result


def test_item_kind_formats_corresponding_custom_label_paragraph():
    original = packet(lines=[('song', 'Music for Reflection\tSelection\tArtist')])
    result = apply_bulletin_rules(original)
    service = result['services'][0]
    assert service['items'][0]['bulletin_text'] == service['bulletin_paragraphs'][0]['text'] == (
        'Music for Reflection\t“Selection”\tArtist')


def test_legacy_display_fields_and_independent_paragraphs_are_formatted():
    original = packet(lines=[('worship_element', 'Hymn\tSelection\tUMH 95')])
    service = original['services'][0]
    item = service['items'][0]
    item['display_wording'] = item.pop('bulletin_text')
    service['bulletin_paragraphs'].append({'style': 'Body', 'text': 'Sermon'})
    before = deepcopy(original)
    result = apply_bulletin_rules(validate_packet(original))
    service = result['services'][0]
    assert service['items'][0]['bulletin_text'] == 'Hymn\t“Selection”\tUMH 95'
    assert service['bulletin_paragraphs'][0]['text'] == service['items'][0]['bulletin_text']
    assert service['bulletin_paragraphs'][1]['text'] == 'Sermon\t\tRev. Nathan Howe'
    assert service['items'][0]['display_wording'] == before['services'][0]['items'][0]['display_wording']
    assert original == before


def test_ambiguous_middle_field_does_not_partially_synchronize_the_sermon():
    original = packet()
    original['services'][0]['items'][0]['bulletin_text'] = 'Sermon\tPossible Title'
    before = deepcopy(original)
    result = apply_bulletin_rules(validate_packet(original))
    assert result['services'] == before['services']
    assert any(i['code'] == 'sermon_speaker_review' for i in result['issues'])


def test_conflicting_evidenced_speakers_are_flagged_without_choosing_one():
    original = packet(lines=[('sermon', 'Sermon'), ('sermon', 'Sermon')])
    for item, name in zip(original['services'][0]['items'], ['Rev. One', 'Rev. Two']):
        item['values'] = {'speaker': {'status': 'supplied', 'value': name,
                                     'evidence': deepcopy(item['evidence'])}}
    before = deepcopy(original)
    result = apply_bulletin_rules(validate_packet(original))
    assert result['services'] == before['services']
    assert any(i['code'] == 'sermon_speaker_review' for i in result['issues'])


@pytest.mark.parametrize('role', ROLES)
def test_build_reads_back_bulletin_and_technical_outputs_and_preserves_edited_handoffs(tmp_path, role):
    original = baseline(tmp_path/'baseline.idml')
    baseline_bytes = original.read_bytes()
    source = tmp_path/'order.txt'
    source.write_bytes(b'Original source evidence')
    data = packet(role, [('prelude', 'Prelude\tSelection\tComposer'), ('sermon', 'Sermon')])
    data['sources'][0]['path'] = str(source)
    data['services'][0]['items'][0]['technical_wording'] = 'Prelude — Selection'
    phase = 'main' if role == 'main' else 'early'
    data['baseline'] = {'path': str(original), 'service': 'main', 'finalized': True,
                        'date': '2099-09-13' if phase == 'main' else data['date'],
                        'sha256': sha256(baseline_bytes).hexdigest()}
    before = deepcopy(data)
    args = (data, tmp_path/'data', {'allowed_styles': ['Body']}, {}, templates(tmp_path))
    result = build_packet(*args, phase=phase)
    idml = next(a for a in result['artifacts'] if a['format'] == 'idml')
    with ZipFile(idml['path']) as output, ZipFile(original) as base:
        texts = [n.text for n in ET.fromstring(output.read('Stories/Story_s1.xml')).iter('Content')]
        assert texts == ['Prelude\t“Selection”\tComposer', 'Sermon\t\tRev. Nathan Howe']
        assert texts[1].split('\t') == ['Sermon', '', 'Rev. Nathan Howe']
        for name in base.namelist():
            if name not in idml['changed_members']:
                assert output.read(name) == base.read(name)
    saved = json.loads((Path(result['manifest_path']).parent/'service-packet.json').read_text(encoding='utf-8'))
    assert saved['services'][0]['items'][1]['bulletin_text'] == texts[1]
    for artifact in result['artifacts']:
        if artifact.get('kind') in {'cameras', 'sound'}:
            column = 0 if artifact['kind'] == 'cameras' else 1
            rows = Document(artifact['path']).tables[0].rows[1:]
            assert [row.cells[column].text for row in rows] == ['Prelude — Selection', 'Sermon']
    assert build_packet(*args, phase=phase)['reused']
    edited = Path(idml['path'])
    edited.write_bytes(b'User-finished handoff')
    later = build_packet(*args, phase=phase)
    assert later['revision'] == result['revision']
    assert any(i['code'] == 'manual_edits_preserved' for i in later['issues'])
    assert edited.read_bytes() == b'User-finished handoff'
    assert original.read_bytes() == baseline_bytes
    assert source.read_bytes() == b'Original source evidence'
    assert data == before
