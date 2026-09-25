from docx import Document
import pytest

from service_packets.technical import render_technical


def templates(tmp_path):
    paths = {}
    for kind, columns in [('cameras', ['Worship Element', 'CAM 1', 'Balcony C', 'Balcony L', 'Proclaim']),
                          ('sound', ['Scene', 'Worship Element'])]:
        doc = Document()
        doc.add_paragraph('old date 10:30 Worship Order')
        table = doc.add_table(rows=2, cols=len(columns))
        for cell, text in zip(table.rows[0].cells, columns):
            cell.text = text
        for cell in table.rows[1].cells:
            cell.text = 'Stale organ cue'
        paths[kind] = tmp_path/f'{kind}-template.docx'
        doc.save(paths[kind])
    return paths


def test_actual_time_dynamic_order_and_unknown_band_cues(tmp_path):
    service = {'key': 'main', 'time': '10:00', 'date': '2099-09-20', 'items': [
        {'id': 'a', 'kind': 'song', 'source_wording': 'Opening song'},
        {'id': 'b', 'kind': 'offertory', 'source_wording': 'Offertory: Simple Kingdom'}]}
    out = tmp_path/'output'
    out.mkdir()
    results = render_technical(service, {}, templates(tmp_path), out)
    assert len(results) == 2
    for result in results:
        document = Document(result['path'])
        assert '10:00' in document.paragraphs[0].text
        assert '10:30' not in document.paragraphs[0].text
        rows = document.tables[0].rows
        assert len(rows) == 3
        column = 0 if result['kind'] == 'cameras' else 1
        assert [row.cells[column].text for row in rows[1:]] == ['Opening song', 'Offertory: Simple Kingdom']
        assert 'REVIEW' in ' '.join(c.text for c in rows[2].cells)
        assert result['item_ids'] == ['a', 'b']


def test_no_early_service_technical_files(tmp_path):
    assert render_technical({'key': 'modern'}, {}, {}, tmp_path) == []


def test_existing_output_is_not_overwritten(tmp_path):
    service = {'key': 'main', 'time': '10:00', 'date': '2099-09-20', 'items': []}
    paths = templates(tmp_path)
    out = tmp_path/'output'
    out.mkdir()
    target = out/'1000-cameras.docx'
    target.write_bytes(b'user edits')
    with pytest.raises(FileExistsError):
        render_technical(service, {}, paths, out)
    assert target.read_bytes() == b'user edits'


def test_technical_spacing_does_not_reuse_bulletin_tabs(tmp_path):
    from docx.oxml.ns import qn
    service = {'key': 'main', 'time': '10:00', 'date': '2099-09-20', 'items': [
        {'id': 'a', 'source_wording': 'Prelude', 'display_wording': 'Prelude\t“Music”\tComposer'}]}
    out = tmp_path/'out'
    out.mkdir()
    result = render_technical(service, {}, templates(tmp_path), out)
    for artifact in result:
        doc = Document(artifact['path'])
        column = 0 if artifact['kind'] == 'cameras' else 1
        assert doc.tables[0].rows[1].cells[column].text == 'Prelude — “Music” — Composer'
        assert doc.sections[0]._sectPr.find(qn('w:vAlign')).get(qn('w:val')) == 'top'


def test_standing_defaults_apply_except_for_explicit_exceptions(tmp_path):
    service = {'key': 'main', 'time': '10:30', 'date': '2099-09-20', 'items': [
        {'id': 'routine', 'source_wording': 'Prelude', 'cue_key': 'organ'},
        {'id': 'special', 'source_wording': 'Solo offertory', 'cue_key': 'organ',
         'cue_exception': 'Soloist replaces organ'}]}
    cues = {'defaults': {'organ': {'convention_id': 'approved-template-organ',
                                 'cameras': ['Organ', '', '', ''], 'sound': ['Organ']}}}
    out = tmp_path/'out'
    out.mkdir()
    for result in render_technical(service, cues, templates(tmp_path), out):
        doc = Document(result['path'])
        cue_column = 1 if result['kind'] == 'cameras' else 0
        assert doc.tables[0].rows[1].cells[cue_column].text == 'Organ'
        assert 'Soloist replaces organ' in doc.tables[0].rows[2].cells[cue_column].text
        assert [finding['item_id'] for finding in result['findings']] == ['special']


def test_concise_technical_wording_for_both_sheets_preserves_cues(tmp_path):
    item = {'id': 'anthem', 'source_wording': 'Anthem',
            'display_wording': 'Anthem\tMusic\tComposer\nSoloist; slide details',
            'technical_wording': 'Anthem\tMusic', 'cue_key': 'choir'}
    service = {'key': 'main', 'time': '10:30', 'date': '2099-09-20', 'items': [item]}
    cues = {'defaults': {'choir': {'convention_id': 'approved',
                                  'cameras': ['Choir', '', '', 'Lyrics'], 'sound': ['Choir']}}}
    out = tmp_path/'out'
    out.mkdir()
    for result in render_technical(service, cues, templates(tmp_path), out):
        row = Document(result['path']).tables[0].rows[1]
        text_column = 0 if result['kind'] == 'cameras' else 1
        cue_column = 1 if result['kind'] == 'cameras' else 0
        assert row.cells[text_column].text == 'Anthem — Music'
        assert row.cells[cue_column].text == 'Choir'
    assert 'Composer' in item['display_wording']


@pytest.mark.parametrize('exit_kind,exit_wording', [('worship_element', 'Exit Music'), ('exit_music', 'Organ recessional')])
def test_camera_stops_before_exit_music_but_sound_retains_full_order(tmp_path, exit_kind, exit_wording):
    items = [{'id': 'response', 'source_wording': 'Benediction Response'},
             {'id': 'postlude', 'source_wording': 'Postlude'},
             {'id': 'chimes', 'source_wording': 'Trinity Chimes'},
             {'id': 'exit', 'source_wording': exit_wording, 'kind': exit_kind}]
    service = {'key': 'main', 'time': '10:30', 'date': '2099-09-20', 'items': items}
    out = tmp_path/'out'
    out.mkdir()
    results = {r['kind']: r for r in render_technical(service, {}, templates(tmp_path), out)}
    camera = results['cameras']
    sound = results['sound']
    assert camera['item_ids'] == ['response', 'postlude', 'chimes']
    assert camera['omitted_item_ids'] == ['exit']
    assert [r.cells[0].text for r in Document(camera['path']).tables[0].rows[1:]] == ['Benediction Response', 'Postlude', 'Trinity Chimes']
    assert sound['item_ids'] == ['response', 'postlude', 'chimes', 'exit']
    assert Document(sound['path']).tables[0].rows[-1].cells[1].text == exit_wording
    assert [i['id'] for i in service['items']] == sound['item_ids']
