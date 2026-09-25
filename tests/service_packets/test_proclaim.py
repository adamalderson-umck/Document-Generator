from docx import Document
import pytest
from service_packets.proclaim import render_proclaim


def test_detailed_fields_and_response_order_are_preserved(tmp_path):
    items = [
        {'id': 'prayer', 'source_wording': 'Pastoral Prayer and Lord’s Prayer',
         'proclaim': [{'label': 'Slide', 'text': 'Vineyard.jpg'}]},
        {'id': 'benediction', 'source_wording': 'Benediction',
         'proclaim': [{'label': 'Text', 'text': 'Go in peace.'}]},
        {'id': 'response', 'source_wording': 'Benediction Response',
         'technical_wording': 'Benediction Response — The Lord Bless You',
         'proclaim': [{'label': 'Composer', 'text': 'Peter C. Lutkin'},
                      {'label': 'Participants', 'text': 'Sanctuary Choir'},
                      {'label': 'Verses', 'text': '1 only'},
                      {'label': 'Lyrics', 'text': 'First line\nSecond line'}]}]
    service = {'key': 'main', 'time': '10:30', 'items': items}
    path = tmp_path/'proclaim.docx'
    result = render_proclaim(service, {'date': '2026-09-20'}, path)
    paragraphs = [p.text for p in Document(path).paragraphs]
    headings = [p.text for p in Document(path).paragraphs if p.style.name == 'Heading 1']
    assert headings == ['01  Pastoral Prayer and Lord’s Prayer', '02  Benediction',
                        '03  Benediction Response — The Lord Bless You']
    for expected in ['Composer: Peter C. Lutkin', 'Participants: Sanctuary Choir',
                     'Verses: 1 only', 'Lyrics: First line\nSecond line', 'Slide: Vineyard.jpg']:
        assert expected in paragraphs
    assert not result['findings']
    before = path.read_bytes()
    with pytest.raises(FileExistsError):
        render_proclaim(service, {'date': '2026-09-20'}, path)
    assert path.read_bytes() == before


def test_missing_enrichment_is_visible_not_silently_complete(tmp_path):
    service = {'key': 'main', 'time': '10:00', 'items': [
        {'id': 'reading', 'source_wording': 'Reading',
         'values': {'reference': {'status': 'supplied', 'value': 'Matthew 20.1–16 (CEB)'}}}]}
    result = render_proclaim(service, {'date': '2026-09-20'}, tmp_path/'p.docx')
    assert result['findings'] == [{'code': 'proclaim_detail_review', 'item_id': 'reading'}]
    text = '\n'.join(p.text for p in Document(result['path']).paragraphs)
    assert 'Matthew 20.1–16 (CEB)' in text and 'REVIEW' in text
