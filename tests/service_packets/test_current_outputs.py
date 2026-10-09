from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from docx import Document
import pytest

from service_packets.build import build_packet
from test_finalize import final_idml


@pytest.fixture
def setup(tmp_path):
    baseline = final_idml(tmp_path/'baseline.idml')
    services = [dict(key=key, time=time, order_source='order', items=[],
                     heading_paragraphs=[dict(style='OOW Body', text=time)],
                     bulletin_paragraphs=[dict(style='OOW Body', text='Original order')])
                for key, time in [('early_traditional', '08:30'), ('modern', '09:30'), ('main', '10:30')]]
    packet = dict(schema_version=1, date='2026-09-13', timezone='America/New_York',
                  schedule=dict(kind='ordinary'), expected_services=[s['key'] for s in services],
                  services=services, sources=[dict(id='order', kind='nathan_docx', scope=[s['key'] for s in services])],
                  baseline=dict(path=str(baseline), sha256=sha256(baseline.read_bytes()).hexdigest(),
                                date='2026-09-06', service='main', finalized=True))
    templates = {}
    for kind, columns in [('cameras', 5), ('sound', 2)]:
        document = Document()
        document.add_paragraph('Header')
        document.add_table(rows=2, cols=columns)
        templates[kind] = tmp_path/(kind+'.docx')
        document.save(templates[kind])
    return packet, tmp_path/'data', {'allowed_styles': ['OOW Body']}, {}, templates


def test_repeated_corrections_leave_exactly_six_outputs_and_preserve_edits(setup):
    packet, root, layout, cues, templates = setup
    main = build_packet(*setup)
    camera = Path(next(a['path'] for a in main['artifacts'] if a.get('kind') == 'cameras'))
    document = Document(camera)
    document.add_paragraph('My saved camera instructions')
    document.save(camera)
    edited = camera.read_bytes()
    early_packet = deepcopy(packet)
    early_packet['baseline']['date'] = packet['date']
    build_packet(early_packet, root, layout, cues, templates, phase='early')
    for n in range(3):
        packet['services'][-1]['bulletin_paragraphs'][0]['text'] = f'Correction {n}'
        result = build_packet(*setup)
        assert camera.read_bytes() == edited
        assert any(i['code'] == 'manual_edits_preserved' for i in result['issues'])
    output = root/'outputs'/packet['date']
    assert sorted(p.name for p in output.iterdir()) == [
        'early_traditional-bulletin.idml', 'main-bulletin.idml', 'main-cameras.docx',
        'main-proclaim.docx', 'main-sound.docx', 'modern-bulletin.idml']
    with ZipFile(output/'main-bulletin.idml') as archive:
        assert b'Correction 2' in archive.read('Stories/s1.xml')


def test_unknown_existing_document_never_becomes_generator_owned(setup):
    packet, root, *_ = setup
    output = root/'outputs'/packet['date']
    output.mkdir(parents=True)
    unknown = output/'main-bulletin.idml'
    unknown.write_bytes(b'Existing file without an ownership record')
    for _ in range(2):
        result = build_packet(*setup)
        assert unknown.read_bytes() == b'Existing file without an ownership record'
        assert any(i['code'] == 'manual_edits_preserved' for i in result['issues'])
    assert len(list(output.iterdir())) == 4


def test_failed_replacement_keeps_last_usable_output(setup, monkeypatch):
    result = build_packet(*setup)
    bulletin = Path(result['artifacts'][0]['path'])
    content = bulletin.read_bytes()
    setup[0]['services'][-1]['bulletin_paragraphs'][0]['text'] = 'New order'
    def fail(*args):
        raise ValueError('Invalid new layout')
    monkeypatch.setattr('service_packets.build.render_idml', fail)
    result = build_packet(*setup)
    assert bulletin.read_bytes() == content
    assert any(i['code'] == 'previous_result_retained' for i in result['issues'])
    assert len(list(bulletin.parent.iterdir())) == 4


def test_edits_made_during_rendering_are_preserved(setup, monkeypatch):
    result = build_packet(*setup)
    bulletin = Path(result['artifacts'][0]['path'])
    setup[0]['services'][-1]['bulletin_paragraphs'][0]['text'] = 'New order'
    from service_packets.build import render_idml
    def edit_while_rendering(*args):
        bulletin.write_bytes(b'User saved during rendering')
        return render_idml(*args)
    monkeypatch.setattr('service_packets.build.render_idml', edit_while_rendering)
    build_packet(*setup)
    assert bulletin.read_bytes() == b'User saved during rendering'
