from pathlib import Path
import pytest

from service_packets.build import build_packet


def test_missing_baseline_does_not_block_technical_outputs(tmp_path):
    from docx import Document
    templates = {}
    for kind, count in [('cameras', 5), ('sound', 2)]:
        doc = Document()
        doc.add_paragraph('Header')
        doc.add_table(rows=2, cols=count)
        path = tmp_path/f'{kind}.docx'
        doc.save(path)
        templates[kind] = path
    packet = {'schema_version': 1, 'date': '2099-09-20', 'timezone': 'America/New_York',
              'schedule': {'kind': 'combined', 'evidence': [{'source_id': 'order', 'location': 'body/p/0'}]},
              'expected_services': ['main'], 'sources': [{'id': 'order', 'kind': 'nathan_docx', 'scope': ['main']}],
              'services': [{'key': 'main', 'time': '10:00', 'title': 'Combined', 'order_source': 'order',
                            'items': [{'id': 'p', 'kind': 'prayer', 'source_wording': 'Opening Prayer',
                                       'evidence': [{'source_id': 'order', 'location': 'body/p/1'}]}]}],
              'baseline': None, 'issues': []}
    first = build_packet(packet, tmp_path/'data', {}, {}, templates)
    second = build_packet(packet, tmp_path/'data', {}, {}, templates)
    assert first['revision'] != second['revision']
    assert len(first['artifacts']) == 3
    assert any(i['code'] == 'baseline_missing' for i in first['issues'])
    assert all(Path(a['path']).is_file() for a in first['artifacts'])
    assert 'pending' in Path(first['review_path']).read_text()
    lock = tmp_path/'data/state/service_packets/build-2099-09-20.lock'
    lock.write_text('Other run still owns the week')
    with pytest.raises(FileExistsError):
        build_packet(packet, tmp_path/'data', {}, {}, templates)
    assert lock.read_text() == 'Other run still owns the week'


@pytest.mark.parametrize('baseline_date,finalized,exception,expected', [
    ('2026-09-20', True, False, 2),
    ('2026-09-13', True, True, 0),
    ('2026-09-27', True, False, 0),
    ('2026-09-20', False, False, 0),
])
def test_early_phase_requires_same_week_final_and_preserves_common_members(
        tmp_path, baseline_date, finalized, exception, expected):
    import json
    from hashlib import sha256
    from zipfile import ZipFile
    from test_idml import baseline
    from service_packets.model import ROLES
    original = baseline(tmp_path/'base.idml')
    services = [{'key': role, 'time': time, 'order_source': 'order', 'items': [],
                 'heading_paragraphs': [{'style': 'Body', 'text': time}],
                 'bulletin_paragraphs': [{'style': 'Body', 'text': role}]}
                for role, time in zip(ROLES, ['08:30', '09:30', '10:30'])]
    packet = {'schema_version': 1, 'date': '2026-09-20', 'timezone': 'America/New_York',
              'schedule': {'kind': 'ordinary'}, 'expected_services': list(ROLES),
              'sources': [{'id': 'order', 'kind': 'nathan_docx', 'scope': list(ROLES)}],
              'services': services, 'baseline': {'path': str(original), 'date': baseline_date,
              'service': 'main', 'finalized': finalized, 'sha256': sha256(original.read_bytes()).hexdigest(),
              'exception_evidence': 'previous-week exception' if exception else None}}
    layout = {'allowed_styles': ['Body']}
    result = build_packet(packet, tmp_path/'data', layout, {}, {}, phase='early')
    assert len(result['artifacts']) == expected
    if not expected:
        assert any(i['code'] == 'baseline_unavailable' for i in result['issues'])
        return
    assert {a['service'] for a in result['artifacts']} == {'early_traditional', 'modern'}
    personnel = [i for i in result['issues'] if i['code'] == 'inherited_personnel_manual_replacement']
    assert {i['service'] for i in personnel} == {'early_traditional', 'modern'}
    assert 'NOT verified' in Path(result['review_path']).read_text(encoding='utf-8')
    for artifact in result['artifacts']:
        with ZipFile(original) as a, ZipFile(artifact['path']) as b:
            for name in a.namelist():
                if name not in ('Stories/Story_s0.xml', 'Stories/Story_s1.xml'):
                    assert a.read(name) == b.read(name)
    repeated = build_packet(packet, tmp_path/'data', layout, {}, {}, phase='early')
    assert repeated['reused'] and repeated['revision'] == result['revision']
    edited = Path(result['artifacts'][0]['path'])
    edited.write_bytes(b'manual artist credit and layout edits')
    next_run = build_packet(packet, tmp_path/'data', layout, {}, {}, phase='early')
    assert next_run['revision'] != result['revision']
    assert edited.read_bytes() == b'manual artist credit and layout edits'


def test_invalid_phase_rejected_before_output(tmp_path):
    with pytest.raises(ValueError, match='phase'):
        build_packet({}, tmp_path, {}, {}, {}, phase='all')
    assert not list(tmp_path.iterdir())
