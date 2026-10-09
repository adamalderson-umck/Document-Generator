from pathlib import Path
from docx import Document
import json
import pytest

from service_packets.sources import read_docx, freeze_sources


def test_docx_preserves_block_order_and_emphasis(tmp_path):
    doc = Document()
    doc.add_paragraph('Before')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run('People: Amen').bold = True
    doc.add_paragraph('After')
    path = tmp_path/'source.docx'
    doc.save(path)
    result = read_docx(path)
    assert [b['kind'] for b in result['blocks']] == ['paragraph', 'table', 'paragraph']
    cell = result['blocks'][1]['rows'][0][0][0]
    assert cell['text'] == 'People: Amen'
    assert cell['runs'][0]['bold'] is True


def test_snapshots_preserve_old_bytes(tmp_path):
    original = tmp_path/'source.txt'
    original.write_text('old')
    candidates = [{'id': 'order', 'path': str(original), 'scope': ['main'], 'kind': 'nathan_docx'}]
    first = freeze_sources(tmp_path/'data', '2099-09-20', candidates)
    original.write_text('new')
    second = freeze_sources(tmp_path/'data', '2099-09-20', candidates)
    assert first != second
    manifest = json.loads((first/'source-manifest.json').read_text())
    assert (first/manifest['sources'][0]['path']).read_text() == 'old'


@pytest.mark.parametrize('bad', ['../escape', 'a/b', 'CON'])
def test_rejects_unsafe_source_identifiers(tmp_path, bad):
    with pytest.raises(ValueError):
        freeze_sources(tmp_path, '2099-09-20', [{'id': bad, 'path': 'unused'}])
