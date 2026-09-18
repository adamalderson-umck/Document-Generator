"""Dynamic main-service tables; unknown operational assignments stay visible."""

from copy import deepcopy
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement


def render_technical(service, cues, templates, directory):
    if service['key'] != 'main':
        return []
    directory = Path(directory)
    prefix = service['time'].replace(':', '')
    targets = {kind: directory/f'{prefix}-{kind}.docx' for kind in ('cameras', 'sound')}
    for target in targets.values():
        if target.exists():
            raise FileExistsError(target)
    results = []
    for kind, target in targets.items():
        doc = Document(templates[kind])
        if not doc.tables or len(doc.tables[0].rows) < 2 or not doc.paragraphs:
            raise ValueError('Technical template requires header, table header and prototype row')
        title = doc.paragraphs[0]
        title.text = f"{service['date']} {service['time']} Worship Order and {'Camera' if kind == 'cameras' else 'Sound'} Plot"
        table = doc.tables[0]
        expected = 5 if kind == 'cameras' else 2
        if len(table.columns) != expected:
            raise ValueError('Unexpected technical template columns')
        row_template = deepcopy(table.rows[1]._tr)
        for row in list(table.rows)[1:]:
            table._tbl.remove(row._tr)
        repeat = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        findings = []
        for item in service['items']:
            table._tbl.append(deepcopy(row_template))
            row = table.rows[-1]
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                cell.text = ''
            text = item.get('display_wording') or item['source_wording']
            row.cells[0 if kind == 'cameras' else 1].text = text
            # Only explicit item assignments with provenance are eligible.
            assignment = cues.get(item['id'], {})
            confirmed = bool(assignment.get('evidence') or assignment.get('convention_id'))
            instructions = assignment.get(kind) if confirmed else None
            indices = [1, 2, 3, 4] if kind == 'cameras' else [0]
            if instructions is None:
                row.cells[indices[0]].text = 'REVIEW: cue not confirmed'
                findings.append({'code': 'cue_unknown', 'item_id': item['id']})
            else:
                if len(instructions) != len(indices):
                    raise ValueError('Cue column count mismatch')
                for index, instruction in zip(indices, instructions):
                    row.cells[index].text = instruction
        with target.open('xb') as stream:
            doc.save(stream)
        results.append({'kind': kind, 'format': 'docx', 'path': str(target),
                        'sha256': sha256(target.read_bytes()).hexdigest(),
                        'item_ids': [i['id'] for i in service['items']],
                        'generation_status': 'generated', 'proof_status': 'pending', 'findings': findings})
    return results
