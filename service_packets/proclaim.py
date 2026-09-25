"""Detailed internal Proclaim sheet, in the same canonical order as technical sheets."""
from hashlib import sha256
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn


def render_proclaim(service, packet, target):
    if service['key'] != 'main':
        raise ValueError('Proclaim is main-service only')
    target = Path(target)
    if target.exists():
        raise FileExistsError(target)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1'):
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(5)
    doc.styles['Heading 1'].font.size = Pt(14)
    doc.add_paragraph('Proclaim Data Sheet', 'Title')
    doc.add_paragraph(f"{packet['date']}  {service['time']} main service", 'Subtitle')
    doc.add_paragraph('Prepare the service in the order below. Review unresolved details before publishing. Pagination is for navigation, not a production constraint.')
    findings = []

    def block(label, text):
        paragraph = doc.add_paragraph()
        if label:
            paragraph.add_run(label + ': ').bold = True
        paragraph.add_run(str(text))

    for index, item in enumerate(service['items'], 1):
        heading = item.get('technical_wording') or item.get('display_wording') or item['source_wording']
        doc.add_paragraph(f"{index:02d}  " + heading.replace('\t', ' — '), 'Heading 1')
        details = item.get('proclaim')
        if details is not None:
            for entry in details:
                block(entry.get('label', ''), entry['text'])
        else:
            display = item.get('display_wording')
            if display and display != heading:
                block('Details', display.replace('\t', ' — '))
            for name, value in item.get('values', {}).items():
                status = value.get('status')
                text = value.get('value', value.get('text'))
                if status in {'supplied', 'explicitly_absent'} and text is not None:
                    block(name.replace('_', ' ').capitalize(), text)
                else:
                    block('REVIEW', name.replace('_', ' ') + ': ' + str(status))
            findings.append({'code': 'proclaim_detail_review', 'item_id': item['id']})
            block('REVIEW', 'Detailed Proclaim enrichment has not been explicitly supplied for this item. Check source text, music credits, verses, participants, lyrics and slide directions.')
        if item.get('cue_exception'):
            block('Technical review', item.get('cue_exception_detail', item['cue_exception']))
    notes = service.get('proclaim_notes', [])
    if notes:
        doc.add_paragraph('Review and source notes', 'Heading 1')
        for entry in notes:
            block(entry.get('label', ''), entry['text'])
    for issue in packet.get('issues', []):
        if issue.get('service') in (None, 'main') and issue.get('message'):
            block('Packet review', issue['message'])
    with target.open('xb') as stream:
        doc.save(stream)
    return {'path': str(target), 'format': 'docx', 'kind': 'proclaim',
            'sha256': sha256(target.read_bytes()).hexdigest(),
            'generation_status': 'generated', 'proof_status': 'pending',
            'item_ids': [item['id'] for item in service['items']], 'findings': findings}
